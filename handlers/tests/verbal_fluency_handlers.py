# handlers/tests/verbal_fluency_handlers.py
import asyncio
import logging
import random
import time
import os # Added os for path.exists

from aiogram import Bot, F, Router
from aiogram.enums import ParseMode, ChatType
from aiogram.exceptions import TelegramBadRequest
from aiogram.fsm.context import FSMContext
from aiogram.types import (
    CallbackQuery,
    InlineKeyboardButton,
    InlineKeyboardMarkup,
    Message,
    Chat,
    User,
)

from fsm_states import VerbalFluencyStates, BatteryCycleStates # Added BatteryCycleStates
from settings import (
    ALL_EXPECTED_HEADERS,
    EXCEL_FILENAME,
    VERBAL_FLUENCY_DURATION_S,
    VERBAL_FLUENCY_TASK_POOL,
    VERBAL_FLUENCY_CATEGORY,
    VERBAL_FLUENCY_HEADERS, # Imported
    BASE_HEADERS # Imported
)
from utils.bot_helpers import (
    send_main_action_menu,
    get_active_profile_from_fsm,
    _clear_fsm_and_set_profile, # Added
    _safe_delete_message # Added
)
from handlers.battery_utils import battery_proceed_after_test_completion # MODIFIED IMPORT
from handlers.common_handlers import TEST_REGISTRY, TEST_SEQUENCE # Import for passing to battery util

from keyboards import (
    ACTION_SELECTION_KEYBOARD_RETURNING,
    ACTION_SELECTION_KEYBOARD_NEW, # Though likely not used here directly
)

logger = logging.getLogger(__name__)
router = Router()
IKB = InlineKeyboardButton


async def _verbal_fluency_timer_task(state: FSMContext, bot_instance: Bot):
    data = await state.get_data()
    chat_id = data.get("vf_chat_id")
    task_message_id = data.get("vf_task_message_id")
    task_letter = data.get("vf_task_letter")
    last_displayed_text = ""
    
    trigger_event_for_stop = data.get("vf_trigger_event_for_stop")
    if not trigger_event_for_stop and chat_id: # Create a mock if not available
        mock_user = User(id=bot_instance.id if hasattr(bot_instance, 'id') else 1, is_bot=True, first_name="Bot")
        mock_chat = Chat(id=chat_id, type=ChatType.PRIVATE)
        trigger_event_for_stop = Message(message_id=0, date=int(time.time()), chat=mock_chat, from_user=mock_user, text="mock for timer end")


    if not all([chat_id, task_message_id, task_letter]):
        logger.error("Verbal Fluency timer: Missing critical data from FSM.")
        await _end_verbal_fluency_test(state, bot_instance, interrupted=True, trigger_event=trigger_event_for_stop)
        return

    base_task_text = f"Задание: Назовите как можно больше слов, начинающихся на букву <b>'{task_letter}'</b>.\n"
    stop_button_markup = InlineKeyboardMarkup(inline_keyboard=[[IKB(text="⏹️ Остановить Тест", callback_data="request_test_stop")]])

    try:
        for i in range(VERBAL_FLUENCY_DURATION_S, -1, -1):
            if await state.get_state() != VerbalFluencyStates.collecting_words.state:
                logger.info("Verbal Fluency timer: State changed, aborting timer.")
                return
            current_timer_display = f"Осталось: {i} сек."
            full_message_content = f"{base_task_text}{current_timer_display}\n\nВводите слова."
            if full_message_content != last_displayed_text:
                try:
                    await bot_instance.edit_message_text(text=full_message_content, chat_id=chat_id, message_id=task_message_id, parse_mode=ParseMode.HTML, reply_markup=stop_button_markup)
                    last_displayed_text = full_message_content
                except TelegramBadRequest as e:
                    if "message is not modified" not in str(e).lower(): logger.warning(f"VF timer: edit_message_text (ID: {task_message_id}) failed: {e}.")
            if i == 0: break
            await asyncio.sleep(1)

        if await state.get_state() == VerbalFluencyStates.collecting_words.state: # Time is up
            logger.info("Verbal Fluency timer: Time is up.")
            await _end_verbal_fluency_test(state, bot_instance, interrupted=False, trigger_event=trigger_event_for_stop)
    except asyncio.CancelledError:
        logger.info("Verbal Fluency timer task explicitly cancelled.")
    except Exception as e:
        logger.error(f"Verbal Fluency timer task unexpected error: {e}", exc_info=True)
        await _end_verbal_fluency_test(state, bot_instance, interrupted=True, trigger_event=trigger_event_for_stop)


async def _end_verbal_fluency_test(
    state: FSMContext, bot_instance: Bot, interrupted: bool,
    trigger_event: Message | CallbackQuery | None,
):
    current_fsm_state_str = await state.get_state()
    if not current_fsm_state_str or not current_fsm_state_str.startswith(VerbalFluencyStates.__name__):
        logger.info("VF _end_test: Called but test not active or already ended.")
        return

    logger.info(f"VF: Entering _end_test. Interrupted: {interrupted}")
    data = await state.get_data()
    chat_id = data.get("vf_chat_id")
    task_message_id = data.get("vf_task_message_id")
    timer_task = data.get("vf_timer_task")

    if timer_task and not timer_task.done():
        timer_task.cancel(); await asyncio.sleep(0.1) # Allow cancellation
    await state.update_data(vf_timer_task=None)

    collected_words = data.get("vf_collected_words", set())
    word_count = len(collected_words)
    
    is_battery = data.get("current_battery_phase") is not None # Check if in battery mode via main FSM
    target_sheet = data.get("current_battery_phase") + "_Ц" if is_battery and data.get("current_battery_phase") else None
    
    await save_verbal_fluency_results(state, is_interrupted=interrupted, target_sheet_name=target_sheet)

    if chat_id:
        final_task_msg_text = "Тест завершен."
        if interrupted: final_task_msg_text = f"Тест на вербальную беглость был <b>ПРЕРВАН</b>. Результаты сохранены."
        else: final_task_msg_text = f"Время вышло! Тест на вербальную беглость завершен."
        
        if task_message_id:
            try:
                await bot_instance.edit_message_text(text=final_task_msg_text, chat_id=chat_id, message_id=task_message_id, reply_markup=None, parse_mode=ParseMode.HTML)
            except TelegramBadRequest: # Message might have been deleted or unpinnable
                logger.warning(f"VF _end_test: Failed to edit task msg {task_message_id}, sending new.")
                try: await bot_instance.send_message(chat_id, final_task_msg_text, parse_mode=ParseMode.HTML)
                except Exception as e_send_final: logger.error(f"VF _end_test: Failed to send final task text: {e_send_final}")
            try: await bot_instance.unpin_chat_message(chat_id=chat_id, message_id=task_message_id)
            except TelegramBadRequest: pass # Ignore if not pinned or already unpinned

        summary_text_for_user = ""
        if interrupted: summary_text_for_user = f"Тест на вербальную беглость был <b>ПРЕРВАН</b>.\nСохраненный результат: {word_count} слов(а)."
        else: summary_text_for_user = (f"Время вышло! Тест на вербальную беглость завершен.\n"
                                       f"Я сохранил результат. Количество названных (уникальных) слов: {word_count}.\n"
                                       f"Общее время выполнения: {VERBAL_FLUENCY_DURATION_S} сек.")
        if not is_battery: # Only send summary if not in battery (battery has its own flow)
             try: await bot_instance.send_message(chat_id, summary_text_for_user, parse_mode=ParseMode.HTML)
             except Exception as e_send_res: logger.error(f"VF _end_test: Fail to send result summary msg: {e_send_res}")
        await state.update_data(vf_task_message_id=None)

    await cleanup_verbal_fluency_ui(state, bot_instance, final_text=None) 
    
    profile_to_set = await get_active_profile_from_fsm(state)
    common_status_msg_id = data.get("status_message_id_to_delete_later")
    if common_status_msg_id and chat_id:
        await _safe_delete_message(bot_instance, chat_id, common_status_msg_id, "VF completion common status cleanup")
    
    await _clear_fsm_and_set_profile(state, profile_to_set) # Clears all but profile keys

    if is_battery:
        # Determine the trigger_event Message object for battery_proceed...
        effective_trigger_msg = trigger_event if isinstance(trigger_event, Message) else (trigger_event.message if trigger_event else None)
        if not effective_trigger_msg and chat_id: # Create mock if necessary
            mock_user = User(id=bot_instance.id if hasattr(bot_instance, 'id') else 1, is_bot=True, first_name="Bot")
            mock_chat = Chat(id=chat_id, type=ChatType.PRIVATE)
            effective_trigger_msg = Message(message_id=0,date=int(time.time()),chat=mock_chat,from_user=mock_user,text="mock")
        if effective_trigger_msg:
            await battery_proceed_after_test_completion(state, bot_instance, effective_trigger_msg, TEST_REGISTRY, TEST_SEQUENCE)
        else: logger.error("VF _end_test: Cannot proceed in battery, missing trigger context and chat_id.")
    elif profile_to_set and chat_id:
        effective_trigger_msg = trigger_event if isinstance(trigger_event, Message) else (trigger_event.message if trigger_event else None)
        if effective_trigger_msg:
            await send_main_action_menu(bot_instance, effective_trigger_msg, ACTION_SELECTION_KEYBOARD_RETURNING)
        else: # Fallback if no message context
            await bot_instance.send_message(chat_id, "Тест завершен. Выберите действие:", reply_markup=ACTION_SELECTION_KEYBOARD_RETURNING)
    elif chat_id: # No profile, standalone
        await bot_instance.send_message(chat_id, "Профиль не активен. Пожалуйста, /start для начала.")
    
    logger.info("Verbal Fluency: Exiting _end_verbal_fluency_test.")


async def start_verbal_fluency_test(
    trigger_event: Message | CallbackQuery, state: FSMContext,
    profile: dict, bot_instance: Bot,
):
    logger.info(f"Starting Verbal Fluency Test for UID: {profile.get('unique_id')}")
    msg_ctx = trigger_event.message if isinstance(trigger_event, CallbackQuery) else trigger_event
    chat_id = msg_ctx.chat.id

    if not VERBAL_FLUENCY_TASK_POOL:
        await bot_instance.send_message(chat_id, "Ошибка: Пул заданий для теста пуст. Тест не может быть запущен.")
        logger.error("Verbal Fluency Test: Task pool is empty.")
        await state.set_state(None)
        active_profile = await get_active_profile_from_fsm(state) # Should be `profile`
        await _clear_fsm_and_set_profile(state, active_profile)
        await send_main_action_menu(bot_instance, msg_ctx, ACTION_SELECTION_KEYBOARD_RETURNING if active_profile else ACTION_SELECTION_KEYBOARD_NEW)
        return

    chosen_task = random.choice(VERBAL_FLUENCY_TASK_POOL)
    task_letter = chosen_task["letter"]

    await state.set_state(VerbalFluencyStates.showing_instructions_and_task)
    await state.update_data(
        vf_unique_id_for_test=profile.get("unique_id"), vf_profile_name_for_test=profile.get("name"),
        vf_profile_age_for_test=profile.get("age"), vf_profile_telegram_id_for_test=profile.get("telegram_id"),
        vf_chat_id=chat_id, vf_task_base_category=VERBAL_FLUENCY_CATEGORY, vf_task_letter=task_letter,
        vf_collected_words=set(), vf_timer_task=None, vf_task_message_id=None,
        vf_trigger_event_for_stop=msg_ctx,
    )
    instruction_text = (f"<b>Тест на вербальную беглость</b>\n\n"
                        f"Вам будет дана буква. Ваша задача – назвать как можно больше слов, "
                        f"начинающихся на эту букву.\n"
                        f"На выполнение задания даётся {VERBAL_FLUENCY_DURATION_S} секунд.\n"
                        f"Слова можно писать в одном или нескольких сообщениях. Каждое слово должно быть не менее двух букв.\n\n"
                        f"Нажмите 'Начать', чтобы увидеть букву и запустить таймер.")
    kbd = InlineKeyboardMarkup(inline_keyboard=[[IKB(text="Начать", callback_data="vf_start_test_confirmed")]])
    
    sent_message = await bot_instance.send_message(chat_id, instruction_text, reply_markup=kbd, parse_mode=ParseMode.HTML)
    await state.update_data(vf_task_message_id=sent_message.message_id)


@router.callback_query(F.data == "vf_start_test_confirmed", VerbalFluencyStates.showing_instructions_and_task)
async def handle_verbal_fluency_start_ack(cb: CallbackQuery, state: FSMContext, bot: Bot):
    await cb.answer()
    data = await state.get_data()
    task_msg_id, task_letter, chat_id = data.get("vf_task_message_id"), data.get("vf_task_letter"), data.get("vf_chat_id")

    if not all([task_msg_id, task_letter, chat_id]):
        logger.error("VF: Missing critical data in FSM for start_ack.")
        if chat_id: await bot.send_message(chat_id, "Произошла ошибка. Пожалуйста, /start.")
        await state.clear(); return

    task_text = (f"Задание: Назовите как можно больше слов, начинающихся на букву <b>'{task_letter}'</b>.\n"
                 f"Осталось: {VERBAL_FLUENCY_DURATION_S} сек.\n\nВводите слова.")
    stop_button_markup = InlineKeyboardMarkup(inline_keyboard=[[IKB(text="⏹️ Остановить Тест", callback_data="request_test_stop")]])

    try:
        await bot.edit_message_text(text=task_text, chat_id=chat_id, message_id=task_msg_id, reply_markup=stop_button_markup, parse_mode=ParseMode.HTML)
    except TelegramBadRequest as e:
        logger.error(f"VF: Failed to edit task message {task_msg_id}: {e}. Sending new.")
        new_msg = await bot.send_message(chat_id, task_text, parse_mode=ParseMode.HTML, reply_markup=stop_button_markup)
        await state.update_data(vf_task_message_id=new_msg.message_id); task_msg_id = new_msg.message_id # update for pin
            
    if task_msg_id and chat_id: # task_msg_id might have been updated
        try: await bot.pin_chat_message(chat_id=chat_id, message_id=task_msg_id, disable_notification=True)
        except TelegramBadRequest as pin_e: logger.error(f"VF: Failed to pin task message {task_msg_id}: {pin_e}")

    await state.set_state(VerbalFluencyStates.collecting_words)
    timer_task = asyncio.create_task(_verbal_fluency_timer_task(state, bot))
    await state.update_data(vf_timer_task=timer_task)


@router.message(VerbalFluencyStates.collecting_words, F.text)
async def handle_verbal_fluency_word_input(message: Message, state: FSMContext):
    data = await state.get_data()
    task_letter = data.get("vf_task_letter", "").lower()
    collected_words_set = data.get("vf_collected_words", set())

    if not task_letter: await message.reply("Ошибка: буква для задания не определена."); return

    user_words_raw = message.text.lower().split()
    newly_added_count = 0
    for word in user_words_raw:
        processed_word = ''.join(filter(str.isalpha, word))
        if len(processed_word) >= 2 and processed_word.startswith(task_letter):
            if processed_word not in collected_words_set:
                collected_words_set.add(processed_word); newly_added_count += 1
    if newly_added_count > 0: await state.update_data(vf_collected_words=collected_words_set)


async def save_verbal_fluency_results(state: FSMContext, is_interrupted: bool, target_sheet_name: Optional[str] = None):
    data = await state.get_data()
    # Prefer battery UID if available
    uid = data.get("original_user_uid") or data.get("vf_unique_id_for_test")
    
    # Get profile info (might be from battery context or vf_ context)
    profile_info = await get_active_profile_from_fsm(state) # Check general active profile
    p_name, p_age, p_tgid = "", "", ""

    if profile_info and profile_info.get('unique_id') == uid: # If active profile matches UID
        p_name, p_age, p_tgid = profile_info.get("name"), profile_info.get("age"), profile_info.get("telegram_id")
    else: # Fallback to vf_ specific data if no active profile or mismatch
        p_name = data.get("vf_profile_name_for_test", data.get("active_name"))
        p_age = data.get("vf_profile_age_for_test", data.get("active_age"))
        p_tgid = data.get("vf_profile_telegram_id_for_test", data.get("active_telegram_id"))


    if not uid: logger.error("VF save: UID not found. Cannot save."); return

    letter = data.get("vf_task_letter", "N/A")
    collected_words = data.get("vf_collected_words", set())
    word_count = len(collected_words)
    words_list_str = ", ".join(sorted(list(collected_words))) if collected_words else "Нет слов"
    interrupted_status = "Да" if is_interrupted else "Нет"
    excel_category_display = f"Слова на букву {letter}"

    logger.info(f"VF Save: UID={uid}, Sheet='{target_sheet_name or 'Default Sheet'}', Cat='{excel_category_display}', L='{letter}', Cnt='{word_count}', Intrp='{interrupted_status}'")
    
    from openpyxl import load_workbook # Local import
    try:
        if not os.path.exists(EXCEL_FILENAME): initialize_excel_file()
        wb = load_workbook(EXCEL_FILENAME)
        
        ws = None
        if target_sheet_name and target_sheet_name in wb.sheetnames: ws = wb[target_sheet_name]
        elif wb.sheetnames: ws = wb[wb.sheetnames[0]]; logger.warning(f"VF Save: Target sheet '{target_sheet_name}' not found/specified. Using '{ws.title}'.")
        else: logger.error(f"VF Save: No sheets in '{EXCEL_FILENAME}'."); return

        sheet_headers = [cell.value for cell in ws[1]] if ws.max_row > 0 else []
        if not sheet_headers or "Unique ID" not in sheet_headers:
            logger.error(f"VF Save: Sheet '{ws.title}' lacks headers or 'Unique ID'. Re-initializing.");
            initialize_excel_file(); wb = load_workbook(EXCEL_FILENAME)
            if target_sheet_name and target_sheet_name in wb.sheetnames: ws = wb[target_sheet_name]
            else: ws = wb[wb.sheetnames[0]]
            sheet_headers = [cell.value for cell in ws[1]]
            if not sheet_headers or "Unique ID" not in sheet_headers:
                raise ValueError(f"Sheet '{ws.title}' still lacks 'Unique ID' after re-init.")

        header_to_idx_map = {header: idx for idx, header in enumerate(sheet_headers)}
        uid_col_idx = header_to_idx_map.get("Unique ID")
        
        row_num = -1
        if uid_col_idx is not None:
            for idx, row_vals_tuple in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                if len(row_vals_tuple) > uid_col_idx and str(row_vals_tuple[uid_col_idx]) == str(uid):
                    row_num = idx; break
        
        if row_num == -1:
            new_row_data = [""] * len(sheet_headers)
            if uid_col_idx is not None: new_row_data[uid_col_idx] = uid
            ws.append(new_row_data); row_num = ws.max_row
            logger.info(f"VF Save: New row for UID {uid} in sheet '{ws.title}' at row {row_num}.")

        # Populate BASE_HEADERS & VERBAL_FLUENCY_HEADERS
        for header_key in BASE_HEADERS + VERBAL_FLUENCY_HEADERS:
            if header_key in header_to_idx_map:
                col_idx = header_to_idx_map[header_key]
                current_val_in_sheet = ws.cell(row=row_num, column=col_idx + 1).value
                value_to_write = None

                if header_key == "Name": value_to_write = p_name
                elif header_key == "Age": value_to_write = p_age
                elif header_key == "Telegram ID": value_to_write = p_tgid
                elif header_key == "Unique ID": value_to_write = uid # Should be there if new row
                elif header_key == "VerbalFluency_Category": value_to_write = excel_category_display
                elif header_key == "VerbalFluency_Letter": value_to_write = letter
                elif header_key == "VerbalFluency_WordCount": value_to_write = word_count
                elif header_key == "VerbalFluency_WordsList": value_to_write = words_list_str
                elif header_key == "VerbalFluency_Interrupted": value_to_write = interrupted_status
                
                if value_to_write is not None and (current_val_in_sheet is None or str(current_val_in_sheet) != str(value_to_write)):
                     ws.cell(row=row_num, column=col_idx + 1).value = value_to_write
            else:
                logger.warning(f"VF Save: Header '{header_key}' not found in sheet '{ws.title}'.")
        
        wb.save(EXCEL_FILENAME)
        logger.info(f"VF results for UID {uid} saved to sheet '{ws.title}'.")
    except Exception as e:
        logger.error(f"VF results save error UID {uid}, Sheet '{target_sheet_name or 'default'}': {e}", exc_info=True)
        # Cannot send message to user from here as bot_instance is not available


async def cleanup_verbal_fluency_ui(state: FSMContext, bot_instance: Bot, final_text: str | None = None):
    logger.info(f"VF: Entering cleanup_ui. Final text: '{final_text}'")
    data = await state.get_data()
    chat_id = data.get("vf_chat_id")
    task_message_id = data.get("vf_task_message_id")
    timer_task = data.get("vf_timer_task")

    if timer_task and not timer_task.done():
        timer_task.cancel(); await asyncio.sleep(0.1) # Allow cancellation

    if final_text and chat_id and task_message_id:
        try: await bot_instance.unpin_chat_message(chat_id=chat_id, message_id=task_message_id)
        except: pass # Ignore errors
        try: await bot_instance.edit_message_text(text=final_text, chat_id=chat_id, message_id=task_message_id, reply_markup=None, parse_mode=ParseMode.HTML)
        except: # If edit fails, try sending new
            try: await bot_instance.send_message(chat_id, final_text, parse_mode=ParseMode.HTML)
            except: pass
    elif final_text and chat_id: # No task_message_id
        try: await bot_instance.send_message(chat_id, final_text, parse_mode=ParseMode.HTML)
        except: pass
    elif task_message_id and chat_id: # No final text, just delete
        try: await bot_instance.unpin_chat_message(chat_id=chat_id, message_id=task_message_id)
        except: pass
        try: await bot_instance.delete_message(chat_id=chat_id, message_id=task_message_id)
        except: pass

    fsm_keys_to_remove = [key for key in data if key.startswith("vf_")]
    for key_to_remove in fsm_keys_to_remove: del data[key_to_remove]
    await state.set_data(data)
    logger.info("Verbal Fluency cleanup: VF-specific FSM data cleaned.")
