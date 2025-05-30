# handlers/tests/raven_matrices_handlers.py
import asyncio
import logging
import random
import time
import os

from aiogram import Bot, F, Router
from aiogram.enums import ParseMode, ChatType
from aiogram.exceptions import TelegramBadRequest
from aiogram.fsm.context import FSMContext
from aiogram.types import (
    CallbackQuery,
    InlineKeyboardButton,
    InlineKeyboardMarkup,
    FSInputFile,
    InputMediaPhoto,
    Message,
    Chat,
    User,
)

from fsm_states import RavenMatricesStates, BatteryCycleStates # Added BatteryCycleStates
import settings as app_settings # Using app_settings prefix
from settings import ( 
    ALL_EXPECTED_HEADERS,
    EXCEL_FILENAME,
    RAVEN_NUM_TASKS_TO_PRESENT,
    # RAVEN_ALL_TASK_FILES, # Accessed via app_settings
    RAVEN_BASE_DIR,
    RAVEN_FEEDBACK_DISPLAY_TIME_S,
    RAVEN_MATRICES_HEADERS, # Imported
    BASE_HEADERS # Imported
)
from utils.image_processors import generate_mr_collage # Not used here, but was in original mental_rotation
from utils.bot_helpers import (
    send_main_action_menu,
    get_active_profile_from_fsm,
    _clear_fsm_and_set_profile, # Added
    _safe_delete_message # Added
)
from handlers.battery_utils import battery_proceed_after_test_completion # MODIFIED IMPORT
from handlers.common_handlers import TEST_REGISTRY, TEST_SEQUENCE # Import for passing to battery util
from keyboards import (
    ACTION_SELECTION_KEYBOARD_RETURNING,
    ACTION_SELECTION_KEYBOARD_NEW, 
)

logger = logging.getLogger(__name__)
router = Router()
IKB = InlineKeyboardButton


def _parse_raven_filename(filename: str) -> tuple[str | None, int | None, int | None]:
    try:
        name_part = os.path.splitext(filename)[0]
        parts = name_part.split('_')
        if len(parts) == 3:
            image_id_str, correct_option_str, num_options_str = parts
            correct_option = int(correct_option_str)
            num_options = int(num_options_str)
            if not (1 <= correct_option <= num_options and num_options > 1):
                logger.warning(f"Raven filename {filename} invalid numbers: correct={correct_option}, total={num_options}")
                return None, None, None
            return image_id_str, correct_option, num_options
        else:
            logger.warning(f"Raven filename {filename} unexpected format. Got {len(parts)} parts.")
            return None, None, None
    except ValueError:
        logger.warning(f"Raven filename {filename} parts not convertible to int.")
        return None, None, None
    except Exception as e:
        logger.error(f"Error parsing Raven filename {filename}: {e}")
        return None, None, None


async def _raven_delayed_feedback_revert(chat_id: int, message_id: int, normal_text: str, bot_instance: Bot, state_at_call: FSMContext):
    try:
        await asyncio.sleep(RAVEN_FEEDBACK_DISPLAY_TIME_S)
        current_fsm_data = await state_at_call.get_data()
        if (current_fsm_data.get("raven_feedback_message_id") == message_id and 
            await state_at_call.get_state() is not None and 
            (await state_at_call.get_state()).startswith(RavenMatricesStates.__name__)):
            try:
                await bot_instance.edit_message_text(text=normal_text, chat_id=chat_id, message_id=message_id, parse_mode=None)
            except TelegramBadRequest as e_edit:
                if "message is not modified" not in str(e_edit).lower() and "message to edit not found" not in str(e_edit).lower():
                    logger.warning(f"Raven delayed feedback (msg {message_id}): Edit failed: {e_edit}")
            except Exception as e_gen_edit: logger.error(f"Raven delayed feedback (msg {message_id}): General error on edit: {e_gen_edit}", exc_info=True)
        else: logger.info(f"Raven delayed feedback (msg {message_id}): State/msg_id changed or test ended. Skipping revert.")
    except asyncio.CancelledError: logger.info(f"Raven delayed feedback task for msg {message_id} cancelled.")
    except Exception as e: logger.error(f"Raven delayed feedback (msg {message_id}): Unexpected error in task: {e}", exc_info=True)


async def _display_raven_task(chat_id: int, state: FSMContext, bot_instance: Bot, trigger_msg_context: Optional[Message]=None):
    data = await state.get_data()
    current_iter_idx = data.get("raven_current_iteration_num", 0)
    session_tasks = data.get("raven_session_task_filenames", [])

    if current_iter_idx >= len(session_tasks):
        logger.info("Raven: All tasks displayed. Finishing test via _display_raven_task.")
        await _finish_raven_matrices_test(state, bot_instance, chat_id, is_interrupted=False, error_occurred=False, trigger_msg_context=trigger_msg_context)
        return

    task_filename_only = session_tasks[current_iter_idx]
    task_image_full_path = os.path.join(RAVEN_BASE_DIR, task_filename_only)
    _, correct_option_1_based, num_total_options = _parse_raven_filename(task_filename_only)

    if not os.path.exists(task_image_full_path) or correct_option_1_based is None or num_total_options is None:
        logger.error(f"Raven: Invalid task file or parsing error for: '{task_filename_only}'")
        if chat_id: await bot_instance.send_message(chat_id, f"Ошибка загрузки задания {current_iter_idx + 1}. Тест Матриц Равена прерван.")
        await _finish_raven_matrices_test(state, bot_instance, chat_id, is_interrupted=True, error_occurred=True, trigger_msg_context=trigger_msg_context)
        return

    await state.update_data(raven_correct_option_for_current_task=correct_option_1_based, raven_num_options_for_current_task=num_total_options, raven_current_task_filename=task_filename_only)
    
    buttons_row, buttons_grid = [], []
    buttons_per_row = 3
    if num_total_options == 8: buttons_per_row = 4
    elif num_total_options <= 4 : buttons_per_row = num_total_options # For 2, 3, 4 options
    
    for i in range(1, num_total_options + 1):
        buttons_row.append(IKB(text=str(i), callback_data=f"raven_answer_{i}"))
        if len(buttons_row) == buttons_per_row or i == num_total_options:
            buttons_grid.append(list(buttons_row)); buttons_row.clear()
    buttons_grid.append([IKB(text="⏹️ Остановить Тест", callback_data="request_test_stop")])
    reply_markup = InlineKeyboardMarkup(inline_keyboard=buttons_grid)

    task_message_id = data.get("raven_task_message_id")
    caption_text = f"Задание {current_iter_idx + 1} из {len(session_tasks)}"

    try:
        if task_message_id:
            media = InputMediaPhoto(media=FSInputFile(task_image_full_path), caption=caption_text)
            await bot_instance.edit_message_media(chat_id=chat_id, message_id=task_message_id, media=media, reply_markup=reply_markup)
        else:
            msg = await bot_instance.send_photo(chat_id=chat_id, photo=FSInputFile(task_image_full_path), caption=caption_text, reply_markup=reply_markup)
            await state.update_data(raven_task_message_id=msg.message_id)
    except (TelegramBadRequest, FileNotFoundError) as e:
        logger.error(f"Raven: Error sending/editing task image '{task_filename_only}': {e}", exc_info=True)
        if chat_id: await bot_instance.send_message(chat_id, "Ошибка отображения задания. Тест Матриц Равена прерван.")
        await _finish_raven_matrices_test(state, bot_instance, chat_id, is_interrupted=True, error_occurred=True, trigger_msg_context=trigger_msg_context)
        return

    await state.update_data(raven_current_task_start_time=time.time())
    await state.set_state(RavenMatricesStates.displaying_task_raven)


async def _finish_raven_matrices_test(state: FSMContext, bot_instance: Bot, chat_id: int | None, 
                                       is_interrupted: bool, error_occurred: bool = False, 
                                       called_by_stop_command: bool = False,
                                       trigger_msg_context: Optional[Message]=None):
    current_fsm_state_on_entry = await state.get_state()
    if not current_fsm_state_on_entry or not current_fsm_state_on_entry.startswith(RavenMatricesStates.__name__):
        logger.info("Raven _finish_test: Called but test not active or already finished.")
        if called_by_stop_command and chat_id: 
             await _safe_delete_message(bot_instance, chat_id, (await state.get_data()).get("status_message_id_to_delete_later"), "Raven stop_test common status")
        return

    logger.info(f"Finishing Raven Test. Interrupted: {is_interrupted}, Error: {error_occurred}, Called by Stop: {called_by_stop_command}")
    data = await state.get_data()
    effective_chat_id = data.get("raven_chat_id", chat_id)
    
    revert_task = data.get("raven_current_feedback_revert_task_ref")
    if revert_task and not revert_task.done(): revert_task.cancel(); await asyncio.sleep(0.02)
    
    results = data.get("raven_iteration_results", [])
    total_done = len(results)
    correct_ans = sum(1 for r in results if r.get("is_correct"))
    start_time, end_time = data.get("raven_total_test_start_time"), data.get("raven_total_test_end_time_actual", time.time())
    total_time_s = round(end_time - start_time, 2) if start_time else 0.0
    correct_times = [r["reaction_time_s"] for r in results if r.get("is_correct") and "reaction_time_s" in r]
    avg_rt_s = round(sum(correct_times) / len(correct_times), 2) if correct_times else 0.0
    ind_resp_str = "; ".join([f"З{idx+1}:{'Прав' if r.get('is_correct') else 'Неправ'},{r.get('reaction_time_s',0):.2f}с" for idx, r in enumerate(results)]) or "N/A"

    await state.update_data(raven_final_correct_answers=correct_ans, raven_final_total_test_time_s=total_time_s, raven_final_avg_time_correct_s=avg_rt_s, raven_final_individual_times_s_str=ind_resp_str, raven_final_interrupted_status=(is_interrupted or error_occurred), raven_final_total_tasks_attempted=total_done)
    
    final_trigger_event = trigger_msg_context or data.get("raven_triggering_event_for_menu")
    if not final_trigger_event and effective_chat_id:
        mock_user = User(id=bot_instance.id if hasattr(bot_instance, 'id') else 1, is_bot=True, first_name="Bot")
        mock_chat = Chat(id=effective_chat_id, type=ChatType.PRIVATE)
        final_trigger_event = Message(message_id=0,date=int(time.time()),chat=mock_chat,from_user=mock_user,text="mock_raven_finish")

    is_battery = data.get("current_battery_phase") is not None
    target_sheet = data.get("current_battery_phase") + "_Ц" if is_battery and data.get("current_battery_phase") else None
    await save_raven_matrices_results(final_trigger_event, state, is_interrupted=(is_interrupted or error_occurred), target_sheet_name=target_sheet)

    if not called_by_stop_command and effective_chat_id:
        num_tasks_in_session = len(data.get("raven_session_task_filenames", []))
        summary_text = "Тест Матриц Равена был прерван."
        if error_occurred: summary_text += " из-за ошибки."
        elif not is_interrupted: summary_text = (f"Тест Матриц Равена успешно завершен!\n"
                                                 f"Правильных ответов: {correct_ans} из {num_tasks_in_session}.\n"
                                                 f"Общее время: {total_time_s:.2f} сек.\n" +
                                                 (f"Среднее время на правильный ответ: {avg_rt_s:.2f} сек." if correct_ans > 0 else "Правильных ответов не было."))
        elif results: summary_text += f"\nЧастичные результаты ({total_done} из {num_tasks_in_session} заданий): {correct_ans} правильных."
        
        if not is_battery: # Only send summary message if not in battery mode
            try: await bot_instance.send_message(effective_chat_id, summary_text, parse_mode=ParseMode.HTML)
            except Exception as e_send_summary: logger.error(f"Raven Finish: Error sending summary: {e_send_summary}")

    await cleanup_raven_ui(state, bot_instance) # Deletes specific Raven messages
    
    if not called_by_stop_command:
        await _safe_delete_message(bot_instance, effective_chat_id, data.get("status_message_id_to_delete_later"), "Raven finish common status")
        profile_to_set = await get_active_profile_from_fsm(state) # Get active profile
        await _clear_fsm_and_set_profile(state, profile_to_set) # Clears all but profile keys
        
        if is_battery:
            if final_trigger_event: await battery_proceed_after_test_completion(state, bot_instance, final_trigger_event, TEST_REGISTRY, TEST_SEQUENCE)
            else: logger.error("Raven Finish: Cannot proceed in battery, missing trigger context.")
        elif profile_to_set and effective_chat_id:
            if final_trigger_event: await send_main_action_menu(bot_instance, final_trigger_event, ACTION_SELECTION_KEYBOARD_RETURNING)
            else: await bot_instance.send_message(effective_chat_id, "Тест завершен. Выберите действие:", reply_markup=ACTION_SELECTION_KEYBOARD_RETURNING)
        elif effective_chat_id:
            await bot_instance.send_message(effective_chat_id, "Тест завершен. Профиль не найден, пожалуйста /start.")


async def start_raven_matrices_test(trigger_event: Message | CallbackQuery, state: FSMContext, profile: dict, bot_instance: Bot):
    logger.info(f"Starting Raven Matrices Test for UID: {profile.get('unique_id', 'N/A')}")
    msg_ctx = trigger_event.message if isinstance(trigger_event, CallbackQuery) else trigger_event
    chat_id = msg_ctx.chat.id

    if not app_settings.RAVEN_ALL_TASK_FILES: # Use app_settings prefix
        logger.error("Raven Start: RAVEN_ALL_TASK_FILES is empty. Cannot start test.")
        await bot_instance.send_message(chat_id, "Ошибка конфигурации: Файлы для Теста Матриц Равена не загружены.")
        await _safe_delete_message(bot_instance, chat_id, (await state.get_data()).get("status_message_id_to_delete_later"), "Raven start common status error")
        await state.clear(); return

    num_tasks = min(RAVEN_NUM_TASKS_TO_PRESENT, len(app_settings.RAVEN_ALL_TASK_FILES))
    session_tasks = random.sample(app_settings.RAVEN_ALL_TASK_FILES, num_tasks)

    await state.set_state(RavenMatricesStates.initial_instructions_raven)
    await state.update_data(
        raven_unique_id_for_test=profile.get("unique_id"), raven_profile_name_for_test=profile.get("name"),
        raven_profile_age_for_test=profile.get("age"), raven_profile_telegram_id_for_test=profile.get("telegram_id"),
        raven_chat_id=chat_id, raven_session_task_filenames=session_tasks, raven_current_iteration_num=0,
        raven_iteration_results=[], raven_total_test_start_time=None, raven_current_task_start_time=None,
        raven_task_message_id=None, raven_feedback_message_id=None, raven_current_feedback_revert_task_ref=None,
        raven_triggering_event_for_menu=msg_ctx,
    )
    instruction_text = (f"<b>Тест Прогрессивных Матриц Равена</b>\n\n"
                        f"Вам будет показана матрица с пропущенным элементом и несколько вариантов для его заполнения. "
                        f"Ваша задача - выбрать наиболее подходящий вариант, чтобы завершить матрицу, следуя логике ее построения.\n\n"
                        f"Тест состоит из {num_tasks} заданий. "
                        f"Постарайтесь отвечать не только правильно, но и быстро. Время ответа учитывается.")
    kbd = InlineKeyboardMarkup(inline_keyboard=[[IKB(text="Начать тест", callback_data="raven_ack_instructions")]])
    
    try:
        await bot_instance.send_message(chat_id, instruction_text, reply_markup=kbd, parse_mode=ParseMode.HTML)
        if isinstance(trigger_event, CallbackQuery) and trigger_event.message: await trigger_event.message.delete()
    except Exception as e_start_instr_send:
        logger.error(f"Raven Start: Error sending initial instructions: {e_start_instr_send}", exc_info=True)
        await bot_instance.send_message(chat_id, "Ошибка при запуске Теста Матриц Равена. Попробуйте /start.")
        await _safe_delete_message(bot_instance, chat_id, (await state.get_data()).get("status_message_id_to_delete_later"), "Raven start common status fail instr")
        await state.clear()


@router.callback_query(F.data == "raven_ack_instructions", RavenMatricesStates.initial_instructions_raven)
async def raven_ack_instructions_callback(cb: CallbackQuery, state: FSMContext, bot: Bot):
    await cb.answer()
    await state.update_data(raven_total_test_start_time=time.time())
    if cb.message: await cb.message.delete()
    chat_id_ack = (await state.get_data()).get("raven_chat_id")
    if chat_id_ack: await _display_raven_task(chat_id_ack, state, bot, trigger_msg_context=cb.message)
    else: 
        logger.error("Raven Ack Instr: chat_id missing. Aborting."); 
        if cb.message: await cb.message.answer("Критическая ошибка: ID чата не найден. Тест прерван.")
        await _finish_raven_matrices_test(state, bot, None, True, True, trigger_msg_context=cb.message)


@router.callback_query(F.data.startswith("raven_answer_"), RavenMatricesStates.displaying_task_raven)
async def handle_raven_answer_callback(cb: CallbackQuery, state: FSMContext, bot: Bot):
    await cb.answer()
    data = await state.get_data()
    chat_id = data.get("raven_chat_id")
    if not chat_id: 
        logger.error("Raven Answer: chat_id missing. Aborting."); 
        if cb.message: await cb.message.answer("Критическая ошибка: ID чата не найден. Тест прерван.")
        await _finish_raven_matrices_test(state, bot, None, True, True, trigger_msg_context=cb.message); return

    task_start_time = data.get("raven_current_task_start_time", time.time())
    reaction_time_s = round(time.time() - task_start_time, 2)
    user_choice_num_1_based = int(cb.data.split("raven_answer_")[-1])
    correct_option_1_based = data.get("raven_correct_option_for_current_task")
    is_correct = user_choice_num_1_based == correct_option_1_based
    
    iteration_result_data = {"task_filename": data.get("raven_current_task_filename", "N/A"), "user_choice": user_choice_num_1_based, "correct_answer_number": correct_option_1_based, "is_correct": is_correct, "reaction_time_s": reaction_time_s}
    current_results_list = data.get("raven_iteration_results", []); current_results_list.append(iteration_result_data)
    await state.update_data(raven_iteration_results=current_results_list)

    feedback_text_bold, feedback_text_normal = f"<b>{'Верно! ✅' if is_correct else 'Неверно!'}</b>", f"{'Верно!' if is_correct else 'Неверно!'}"
    feedback_msg_id = data.get("raven_feedback_message_id")
    
    previous_revert_task = data.get("raven_current_feedback_revert_task_ref")
    if previous_revert_task and not previous_revert_task.done(): previous_revert_task.cancel(); await asyncio.sleep(0.01)

    try:
        if feedback_msg_id: await bot.edit_message_text(text=feedback_text_bold, chat_id=chat_id, message_id=feedback_msg_id, parse_mode=ParseMode.HTML)
        else: msg_fb = await bot.send_message(chat_id, feedback_text_bold, parse_mode=ParseMode.HTML); feedback_msg_id = msg_fb.message_id; await state.update_data(raven_feedback_message_id=feedback_msg_id)
        if feedback_msg_id:
            revert_task = asyncio.create_task(_raven_delayed_feedback_revert(chat_id, feedback_msg_id, feedback_text_normal, bot, state))
            await state.update_data(raven_current_feedback_revert_task_ref=revert_task)
    except Exception as e_fb_ans_cb: logger.error(f"Raven Answer: Feedback error: {e_fb_ans_cb}", exc_info=True)
    
    current_iter_idx, session_tasks_len = data.get("raven_current_iteration_num", 0), len(data.get("raven_session_task_filenames", []))
    next_iter_idx = current_iter_idx + 1
    await state.update_data(raven_current_iteration_num=next_iter_idx)

    if next_iter_idx < session_tasks_len: await _display_raven_task(chat_id, state, bot, trigger_msg_context=cb.message)
    else: 
        await state.update_data(raven_total_test_end_time_actual=time.time())
        logger.info("Raven Matrices Test: All iterations completed by user.")
        await _finish_raven_matrices_test(state, bot, chat_id, is_interrupted=False, error_occurred=False, trigger_msg_context=cb.message)


async def save_raven_matrices_results(
    trigger_msg_context: Optional[Message], state: FSMContext, 
    is_interrupted: bool = False, target_sheet_name: Optional[str] = None,
):
    logger.info(f"Saving Raven results. Interrupted: {is_interrupted}. Sheet: {target_sheet_name or 'default'}")
    data = await state.get_data()
    uid = data.get("original_user_uid") or data.get("raven_unique_id_for_test") or data.get("active_unique_id")
    
    profile_info = await get_active_profile_from_fsm(state)
    p_name, p_age, p_tgid = "", "", ""
    if profile_info and profile_info.get('unique_id') == uid:
        p_name, p_age, p_tgid = profile_info.get("name"), profile_info.get("age"), profile_info.get("telegram_id")
    else:
        p_name, p_age, p_tgid = data.get("raven_profile_name_for_test"), data.get("raven_profile_age_for_test"), data.get("raven_profile_telegram_id_for_test")

    if not uid: logger.error("Raven Save: UID not found."); return

    correct_ans = data.get("raven_final_correct_answers", 0)
    total_time = data.get("raven_final_total_test_time_s", 0.0)
    avg_rt_correct = data.get("raven_final_avg_time_correct_s", 0.0)
    ind_times_str = data.get("raven_final_individual_times_s_str", "N/A")
    interrupted_status = "Да" if data.get("raven_final_interrupted_status", is_interrupted) else "Нет"

    from openpyxl import load_workbook # Local import
    try:
        if not os.path.exists(EXCEL_FILENAME): initialize_excel_file()
        wb = load_workbook(EXCEL_FILENAME)
        
        ws = None
        if target_sheet_name and target_sheet_name in wb.sheetnames: ws = wb[target_sheet_name]
        elif wb.sheetnames: ws = wb[wb.sheetnames[0]]; logger.warning(f"Raven Save: Target sheet '{target_sheet_name}' not found/specified. Using '{ws.title}'.")
        else: logger.error(f"Raven Save: No sheets in '{EXCEL_FILENAME}'."); return

        sheet_headers = [cell.value for cell in ws[1]] if ws.max_row > 0 else []
        if not sheet_headers or "Unique ID" not in sheet_headers:
            logger.error(f"Raven Save: Sheet '{ws.title}' lacks headers or 'Unique ID'. Re-initializing.");
            initialize_excel_file(); wb = load_workbook(EXCEL_FILENAME)
            if target_sheet_name and target_sheet_name in wb.sheetnames: ws = wb[target_sheet_name]
            else: ws = wb[wb.sheetnames[0]]
            sheet_headers = [cell.value for cell in ws[1]]
            if not sheet_headers or "Unique ID" not in sheet_headers:
                 raise ValueError(f"Sheet '{ws.title}' still lacks 'Unique ID' after re-init.")
        
        header_to_idx_map = {header: idx for idx, header in enumerate(sheet_headers)}
        uid_col_idx = header_to_idx_map.get("Unique ID")
        
        row_num = -1
        if uid_col_idx is not None:
            for r_idx, row_vals_tuple in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                if len(row_vals_tuple) > uid_col_idx and str(row_vals_tuple[uid_col_idx]) == str(uid):
                    row_num = r_idx; break
        
        if row_num == -1:
            new_row_data = [""] * len(sheet_headers)
            if uid_col_idx is not None: new_row_data[uid_col_idx] = uid
            ws.append(new_row_data); row_num = ws.max_row
            logger.info(f"Raven Save: New row for UID {uid} in sheet '{ws.title}' at row {row_num}.")

        # Populate BASE_HEADERS & RAVEN_MATRICES_HEADERS
        for header_key in BASE_HEADERS + RAVEN_MATRICES_HEADERS:
            if header_key in header_to_idx_map:
                col_idx = header_to_idx_map[header_key]
                current_val_in_sheet = ws.cell(row=row_num, column=col_idx + 1).value
                value_to_write = None

                if header_key == "Name": value_to_write = p_name
                elif header_key == "Age": value_to_write = p_age
                elif header_key == "Telegram ID": value_to_write = p_tgid
                elif header_key == "Unique ID": value_to_write = uid
                elif header_key == "RavenMatrices_CorrectAnswers": value_to_write = correct_ans
                elif header_key == "RavenMatrices_TotalTime_s": value_to_write = total_time
                elif header_key == "RavenMatrices_AvgTimeCorrect_s": value_to_write = avg_rt_correct
                elif header_key == "RavenMatrices_IndividualTimes_s": value_to_write = ind_times_str
                elif header_key == "RavenMatrices_Interrupted": value_to_write = interrupted_status
                
                if value_to_write is not None and (current_val_in_sheet is None or str(current_val_in_sheet) != str(value_to_write)):
                     ws.cell(row=row_num, column=col_idx + 1).value = value_to_write
            else:
                logger.warning(f"Raven Save: Header '{header_key}' not found in sheet '{ws.title}'.")
        
        wb.save(EXCEL_FILENAME)
        logger.info(f"Raven results for UID {uid} saved to sheet '{ws.title}'. Interrupted: {interrupted_status}")
    except Exception as e:
        logger.error(f"Raven Save Results: Excel error for UID {uid}, Sheet '{target_sheet_name or 'default'}': {e}", exc_info=True)
        if trigger_msg_context and hasattr(trigger_msg_context, 'chat') and await state.get_state() is not None:
            try: await trigger_msg_context.answer("Ошибка при сохранении результатов Теста Матриц Равена.")
            except: pass


async def cleanup_raven_ui(state: FSMContext, bot_instance: Bot, final_text: Optional[str] = None):
    data = await state.get_data()
    chat_id = data.get("raven_chat_id")
    logger.info(f"Raven Cleanup UI: Chat {chat_id if chat_id else 'N/A'}. Final text (ignored for task msg edit): '{final_text}'")

    revert_task = data.get("raven_current_feedback_revert_task_ref")
    if revert_task and not revert_task.done(): revert_task.cancel(); await asyncio.sleep(0.02) # give it a moment
    
    task_msg_id = data.get("raven_task_message_id")
    feedback_msg_id = data.get("raven_feedback_message_id")

    if chat_id:
        if task_msg_id:
            try: await bot_instance.delete_message(chat_id, task_msg_id)
            except: pass 
        if feedback_msg_id and feedback_msg_id != task_msg_id: # Avoid double delete if same
            try: await bot_instance.delete_message(chat_id, feedback_msg_id)
            except: pass
        # If final_text is provided (e.g. by stop_test_command_handler), it will send its own message.
        # This cleanup focuses on removing test-specific UI.

    fsm_keys_to_remove = [key for key in data if key.startswith("raven_")]
    for key_to_remove in fsm_keys_to_remove: del data[key_to_remove]
    await state.set_data(data)
    logger.info(f"Raven Cleanup UI: FSM data cleaned of raven_ prefixed keys. Chat: {chat_id or 'N/A'}")
