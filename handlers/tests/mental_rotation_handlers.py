# handlers/tests/mental_rotation_handlers.py
import asyncio
import logging
import random
import time
import os

from aiogram import Bot, F, Router
from aiogram.enums import ParseMode, ChatType
from aiogram.exceptions import TelegramBadRequest
from aiogram.fsm.context import FSMContext
from aiogram.types import (
    CallbackQuery,
    InlineKeyboardButton,
    InlineKeyboardMarkup,
    FSInputFile,
    InputMediaPhoto,
    Message,
    Chat,
    User,
)

import settings as app_settings # Changed to import settings as app_settings
from fsm_states import MentalRotationStates, BatteryCycleStates # Added BatteryCycleStates
# Explicitly import headers needed for save function
from settings import (
    ALL_EXPECTED_HEADERS,
    EXCEL_FILENAME,
    MENTAL_ROTATION_NUM_ITERATIONS,
    MR_REFERENCES_DIR, # This is settings.MR_REFERENCES_DIR
    MR_CORRECT_PROJECTIONS_DIR, # This is settings.MR_CORRECT_PROJECTIONS_DIR
    MR_FEEDBACK_DISPLAY_TIME_S,
    MENTAL_ROTATION_HEADERS, 
    BASE_HEADERS
)
from utils.image_processors import generate_mr_collage
from utils.bot_helpers import (
    send_main_action_menu,
    get_active_profile_from_fsm,
    _clear_fsm_and_set_profile, # Added
    _safe_delete_message # Added
)
from handlers.battery_utils import battery_proceed_after_test_completion # MODIFIED IMPORT
from ..common_handlers import TEST_REGISTRY, BATTERY_TEST_SEQUENCE_KEYS # MODIFIED IMPORT

from keyboards import (
    ACTION_SELECTION_KEYBOARD_RETURNING,
    ACTION_SELECTION_KEYBOARD_NEW, 
)

logger = logging.getLogger(__name__)
router = Router()
IKB = InlineKeyboardButton


async def _get_mr_stimulus_for_iteration(
    state: FSMContext,
) -> tuple[str | None, list[str] | None, int | None, str | None]:
    data = await state.get_data()
    used_references = data.get("mr_used_references", [])

    if not app_settings.MR_REFERENCE_FILES:
        return None, None, None, "Пул эталонных изображений пуст."
    available_references = [ref for ref in app_settings.MR_REFERENCE_FILES if ref not in used_references]
    if not available_references: return None, None, None, "Больше нет уникальных эталонных изображений."

    selected_reference_filename = random.choice(available_references)
    selected_reference_path = os.path.join(MR_REFERENCES_DIR, selected_reference_filename)
    if not os.path.exists(selected_reference_path): return None, None, None, f"Эталонное изображение не найдено: {selected_reference_filename}"
    
    used_references.append(selected_reference_filename)

    if not app_settings.MR_CORRECT_PROJECTIONS_MAP:
        return None, None, None, "Карта правильных проекций пуста."
    correct_projection_filenames = app_settings.MR_CORRECT_PROJECTIONS_MAP.get(selected_reference_filename, [])
    if not correct_projection_filenames: return None, None, None, f"Нет карты правильных проекций для {selected_reference_filename}"
    
    chosen_correct_proj_filename = random.choice(correct_projection_filenames)
    correct_projection_path = os.path.join(MR_CORRECT_PROJECTIONS_DIR, chosen_correct_proj_filename)
    if not os.path.exists(correct_projection_path): return None, None, None, f"Файл правильной проекции не найден: {chosen_correct_proj_filename}"

    if not app_settings.MR_ALL_DISTRACTORS_FILES: return None, None, None, "Пул дистракторов пуст."
    
    num_distractors_to_select = 3
    if len(app_settings.MR_ALL_DISTRACTORS_FILES) < num_distractors_to_select:
        return None, None, None, f"Недостаточно дистракторов (нужно {num_distractors_to_select})."
    
    valid_distractors = [dp for dp in app_settings.MR_ALL_DISTRACTORS_FILES if os.path.exists(dp)]
    if len(valid_distractors) < num_distractors_to_select:
        return None, None, None, f"Недостаточно ВАЛИДНЫХ дистракторов (нужно {num_distractors_to_select})."

    selected_distractor_paths = random.sample(valid_distractors, num_distractors_to_select)
    await state.update_data(mr_used_references=used_references)

    options_paths = [correct_projection_path] + selected_distractor_paths
    random.shuffle(options_paths)
    correct_option_index = options_paths.index(correct_projection_path)
    return selected_reference_path, options_paths, correct_option_index, None


async def _mr_schedule_feedback_revert(chat_id: int, message_id: int, normal_text: str, bot_instance: Bot, state_context_at_call: FSMContext):
    try:
        await asyncio.sleep(MR_FEEDBACK_DISPLAY_TIME_S)
        current_fsm_state_val = await state_context_at_call.get_state()
        current_fsm_data = await state_context_at_call.get_data()
        if current_fsm_state_val is not None and current_fsm_state_val.startswith(MentalRotationStates.__name__) and current_fsm_data.get("mr_feedback_message_id") == message_id:
            try:
                await bot_instance.edit_message_text(text=normal_text, chat_id=chat_id, message_id=message_id, parse_mode=None)
            except TelegramBadRequest as e_edit:
                if "message is not modified" not in str(e_edit).lower() and "message to edit not found" not in str(e_edit).lower():
                    logger.warning(f"MR Feedback Revert (msg {message_id}): Edit failed: {e_edit}")
            except Exception as e_gen_edit: logger.error(f"MR Feedback Revert (msg {message_id}): General error on edit: {e_gen_edit}", exc_info=True)
        else: logger.info(f"MR Feedback Revert (msg {message_id}): State/msg_id changed or test ended. Skipping revert.")
    except asyncio.CancelledError: logger.info(f"MR Feedback Revert task for msg {message_id} was cancelled.")
    except Exception as e: logger.error(f"MR Feedback Revert (msg {message_id}): Unexpected error in task: {e}", exc_info=True)


async def _display_mr_stimulus(chat_id: int, state: FSMContext, bot_instance: Bot, is_editing: bool = False, trigger_msg_context: Optional[Message]=None):
    data = await state.get_data()
    current_iteration = data.get("mr_current_iteration", 0) + 1
    await state.update_data(mr_current_iteration=current_iteration)

    ref_path, opt_paths, correct_idx, err_msg = await _get_mr_stimulus_for_iteration(state)

    if err_msg or not ref_path or not opt_paths or correct_idx is None:
        logger.error(f"MR Display Stimulus: Error from _get_mr_stimulus_for_iteration - {err_msg}")
        if chat_id: await bot_instance.send_message(chat_id, f"Ошибка подготовки задания: {err_msg}. Тест умственного вращения прерван.")
        await _finish_mental_rotation_test(state, bot_instance, chat_id, is_interrupted=True, error_occurred=True, trigger_msg_context=trigger_msg_context)
        return

    collage_input_file, path_to_delete_collage = None, None
    try:
        collage_file_path_or_bytes = await generate_mr_collage(opt_paths)
        if isinstance(collage_file_path_or_bytes, str):
            collage_input_file, path_to_delete_collage = FSInputFile(collage_file_path_or_bytes), collage_file_path_or_bytes
        elif hasattr(collage_file_path_or_bytes, 'read'): collage_input_file = collage_file_path_or_bytes
        else: raise ValueError("generate_mr_collage returned an unexpected type")

        if not collage_input_file: raise ValueError("Collage generation failed.")
        
        await state.update_data(mr_correct_option_index_for_current_iter=correct_idx)
        ref_msg_id, options_msg_id = data.get("mr_reference_message_id"), data.get("mr_options_message_id")

        if not is_editing: # Delete old messages if not editing in place (first iteration)
            for mid_key in ["mr_reference_message_id", "mr_options_message_id", "mr_countdown_message_id", "mr_feedback_message_id"]:
                old_mid = data.get(mid_key)
                if old_mid and chat_id: 
                    try: await bot_instance.delete_message(chat_id, old_mid)
                    except: pass
                await state.update_data({mid_key: None})
            ref_msg_id, options_msg_id = None, None # Ensure they are None for sending new

        if chat_id: # Ensure chat_id is valid before sending/editing
            if is_editing and ref_msg_id: await bot_instance.edit_message_media(chat_id=chat_id, message_id=ref_msg_id, media=InputMediaPhoto(media=FSInputFile(ref_path)))
            else: msg_ref = await bot_instance.send_photo(chat_id, FSInputFile(ref_path)); await state.update_data(mr_reference_message_id=msg_ref.message_id)

            buttons = [[IKB(text="1",callback_data="mr_answer_1"),IKB(text="2",callback_data="mr_answer_2")],[IKB(text="3",callback_data="mr_answer_3"),IKB(text="4",callback_data="mr_answer_4")],[IKB(text="⏹️ Остановить Тест", callback_data="request_test_stop")]]
            reply_markup = InlineKeyboardMarkup(inline_keyboard=buttons)

            if is_editing and options_msg_id: await bot_instance.edit_message_media(chat_id=chat_id, message_id=options_msg_id, media=InputMediaPhoto(media=collage_input_file), reply_markup=reply_markup)
            else: msg_opts = await bot_instance.send_photo(chat_id, collage_input_file, reply_markup=reply_markup); await state.update_data(mr_options_message_id=msg_opts.message_id)
        else: raise ValueError("chat_id is None for MR stimulus display.")

        await state.update_data(mr_iteration_start_time=time.time())
        await state.set_state(MentalRotationStates.displaying_stimulus_mr)
    except Exception as e_stim:
        logger.error(f"MR Display Stimulus: General error: {e_stim}", exc_info=True)
        if chat_id: await bot_instance.send_message(chat_id, "Критическая ошибка при отображении задания. Тест прерван.")
        await _finish_mental_rotation_test(state, bot_instance, chat_id, is_interrupted=True, error_occurred=True, trigger_msg_context=trigger_msg_context)
    finally:
        if path_to_delete_collage and os.path.exists(path_to_delete_collage):
            try: os.remove(path_to_delete_collage)
            except Exception as e_del_collage: logger.error(f"MR Display: Failed to delete temp collage file {path_to_delete_collage}: {e_del_collage}")


async def _mr_proceed_to_next_iteration_or_finish(state: FSMContext, bot_instance: Bot, chat_id: int, trigger_msg_context: Optional[Message]=None):
    data = await state.get_data()
    current_iteration = data.get("mr_current_iteration", 0)
    if current_iteration < MENTAL_ROTATION_NUM_ITERATIONS:
        await state.set_state(MentalRotationStates.inter_iteration_countdown_mr)
        countdown_task = asyncio.create_task(_mr_inter_iteration_countdown_task(state, bot_instance, chat_id, trigger_msg_context))
        await state.update_data(mr_inter_iteration_countdown_task_ref=countdown_task)
    else:
        await _finish_mental_rotation_test(state, bot_instance, chat_id, is_interrupted=False, trigger_msg_context=trigger_msg_context)


async def _mr_inter_iteration_countdown_task(state: FSMContext, bot_instance: Bot, chat_id: int, trigger_msg_context: Optional[Message]=None):
    await asyncio.sleep(0.2) # Brief pause after feedback clear
    if await state.get_state() != MentalRotationStates.inter_iteration_countdown_mr.state: return

    countdown_msg_id_local = None
    try:
        if not chat_id: raise ValueError("chat_id is missing")
        countdown_msg = await bot_instance.send_message(chat_id, "Следующее задание через: 3...")
        countdown_msg_id_local = countdown_msg.message_id
        await state.update_data(mr_countdown_message_id=countdown_msg_id_local)

        for i in range(2, 0, -1):
            await asyncio.sleep(1)
            if await state.get_state() != MentalRotationStates.inter_iteration_countdown_mr.state: return
            await bot_instance.edit_message_text(text=f"Следующее задание через: {i}...", chat_id=chat_id, message_id=countdown_msg_id_local)
        
        await asyncio.sleep(1)
        if await state.get_state() != MentalRotationStates.inter_iteration_countdown_mr.state: return
        
        if countdown_msg_id_local: await bot_instance.delete_message(chat_id, countdown_msg_id_local)
        await state.update_data(mr_countdown_message_id=None)
        await _display_mr_stimulus(chat_id, state, bot_instance, is_editing=True, trigger_msg_context=trigger_msg_context)

    except (TelegramBadRequest, ValueError) as e_tb_val: # Catch expected errors
        logger.warning(f"MR Countdown: Expected error: {e_tb_val}. Attempting recovery/finish.")
        if await state.get_state() == MentalRotationStates.inter_iteration_countdown_mr.state:
            if countdown_msg_id_local and chat_id : 
                try: await bot_instance.delete_message(chat_id, countdown_msg_id_local)
                except: pass
            if chat_id: await _display_mr_stimulus(chat_id, state, bot_instance, is_editing=True, trigger_msg_context=trigger_msg_context) # Try to proceed
            else: await _finish_mental_rotation_test(state, bot_instance, None, True, True, trigger_msg_context=trigger_msg_context)
    except asyncio.CancelledError: logger.info(f"MR Countdown task for chat {chat_id} cancelled.")
    except Exception as e_unexp:
        logger.error(f"MR Countdown: Unexpected error: {e_unexp}", exc_info=True)
        if await state.get_state() == MentalRotationStates.inter_iteration_countdown_mr.state:
            await _finish_mental_rotation_test(state, bot_instance, chat_id, True, True, trigger_msg_context=trigger_msg_context)
    finally: await state.update_data(mr_inter_iteration_countdown_task_ref=None)


async def _finish_mental_rotation_test(state: FSMContext, bot_instance: Bot, chat_id: int | None, 
                                       is_interrupted: bool, error_occurred: bool = False, 
                                       called_by_stop_command: bool = False, 
                                       trigger_msg_context: Optional[Message]=None):
    current_fsm_state_on_entry = await state.get_state()
    if not current_fsm_state_on_entry or not current_fsm_state_on_entry.startswith(MentalRotationStates.__name__):
        logger.info("MR _finish_test: Called but test not in an active MR state or already finished.")
        if called_by_stop_command and chat_id: # From /stoptest
             await _safe_delete_message(bot_instance, chat_id, (await state.get_data()).get("status_message_id_to_delete_later"), "MR stop_test common status")
        return

    logger.info(f"Finishing MR Test. Interrupted: {is_interrupted}, Error: {error_occurred}, Called by Stop: {called_by_stop_command}")
    data = await state.get_data()
    effective_chat_id = data.get("mr_chat_id", chat_id)
    
    # Cancel any running async tasks specific to MR
    for task_key in ["mr_current_feedback_revert_task_ref", "mr_inter_iteration_countdown_task_ref"]:
        task = data.get(task_key)
        if task and not task.done(): task.cancel(); await asyncio.sleep(0.02) 
    
    results, start_time = data.get("mr_iteration_results", []), data.get("mr_test_start_time")
    correct_ans = sum(1 for r in results if r.get("is_correct"))
    total_done = len(results)
    total_time_s = round(time.time() - start_time, 2) if start_time else 0.0
    correct_times = [r["reaction_time_s"] for r in results if r.get("is_correct") and "reaction_time_s" in r]
    avg_rt_s = round(sum(correct_times) / len(correct_times), 2) if correct_times else 0.0
    ind_resp_str = "; ".join([f"И{r.get('iteration','?')}:{'Прав' if r.get('is_correct') else 'Неправ'},{r.get('reaction_time_s',0):.2f}с" for r in results]) or "N/A"

    await state.update_data(mr_final_correct_answers=correct_ans, mr_final_avg_reaction_time_s=avg_rt_s, mr_final_total_test_time_s=total_time_s, mr_final_individual_responses_str=ind_resp_str, mr_final_interrupted_status=(is_interrupted or error_occurred))
    
    # Determine effective trigger_msg_context for saving and menu navigation
    final_trigger_event = trigger_msg_context or data.get("mr_triggering_event_for_menu")
    if not final_trigger_event and effective_chat_id: # Create mock if still none
        mock_user = User(id=bot_instance.id if hasattr(bot_instance, 'id') else 1, is_bot=True, first_name="Bot")
        mock_chat = Chat(id=effective_chat_id, type=ChatType.PRIVATE)
        final_trigger_event = Message(message_id=0,date=int(time.time()),chat=mock_chat,from_user=mock_user,text="mock_mr_finish")
    
    is_battery = data.get("current_battery_phase") is not None
    target_sheet = data.get("current_battery_phase") + "_Ц" if is_battery and data.get("current_battery_phase") else None
    await save_mental_rotation_results(final_trigger_event, state, is_interrupted=(is_interrupted or error_occurred), target_sheet_name=target_sheet)

    if not called_by_stop_command and effective_chat_id:
        summary_text = "Тест умственного вращения был прерван."
        if error_occurred: summary_text += " из-за ошибки."
        elif not is_interrupted: summary_text = (f"Тест умственного вращения успешно завершен!\n"
                                                 f"Правильных ответов: {correct_ans}/{MENTAL_ROTATION_NUM_ITERATIONS}\n"
                                                 f"Среднее время реакции на правильные: {avg_rt_s:.2f} сек\n"
                                                 f"Общее время теста: {total_time_s:.2f} сек")
        if not is_battery: # Only send summary if not in battery
            try: await bot_instance.send_message(effective_chat_id, summary_text, parse_mode=ParseMode.HTML)
            except Exception as e_send_summary: logger.error(f"MR Finish: Error sending summary: {e_send_summary}")

    await cleanup_mental_rotation_ui(state, bot_instance, final_text=None if not called_by_stop_command else "Тест умственного вращения остановлен.")
    
    if not called_by_stop_command: # If not stopped by /stoptest, then handle FSM and menu
        await _safe_delete_message(bot_instance, effective_chat_id, data.get("status_message_id_to_delete_later"), "MR finish common status")
        profile_to_set = await get_active_profile_from_fsm(state)
        await _clear_fsm_and_set_profile(state, profile_to_set)
        
        if is_battery:
            if final_trigger_event: await battery_proceed_after_test_completion(state, bot_instance, final_trigger_event, TEST_REGISTRY, BATTERY_TEST_SEQUENCE_KEYS) # MODIFIED CALL
            else: logger.error("MR Finish: Cannot proceed in battery, missing trigger context.")
        elif profile_to_set and effective_chat_id:
            if final_trigger_event: await send_main_action_menu(bot_instance, final_trigger_event, ACTION_SELECTION_KEYBOARD_RETURNING)
            else: await bot_instance.send_message(effective_chat_id, "Тест завершен. Выберите действие:", reply_markup=ACTION_SELECTION_KEYBOARD_RETURNING)
        elif effective_chat_id: # No profile
            await bot_instance.send_message(effective_chat_id, "Тест завершен. Профиль не найден, пожалуйста /start.")


async def start_mental_rotation_test(trigger_event: Message | CallbackQuery, state: FSMContext, profile: dict, bot_instance: Bot):
    logger.info(f"Starting Mental Rotation Test for UID: {profile.get('unique_id', 'N/A')}")
    msg_ctx = trigger_event.message if isinstance(trigger_event, CallbackQuery) else trigger_event
    chat_id = msg_ctx.chat.id

    await state.set_state(MentalRotationStates.initial_instructions_mr)
    await state.update_data(
        mr_unique_id_for_test=profile.get("unique_id"), mr_profile_name_for_test=profile.get("name"),
        mr_profile_age_for_test=profile.get("age"), mr_profile_telegram_id_for_test=profile.get("telegram_id"),
        mr_chat_id=chat_id, mr_current_iteration=0, mr_iteration_results=[], mr_used_references=[],
        mr_test_start_time=None, mr_reference_message_id=None, mr_options_message_id=None,
        mr_countdown_message_id=None, mr_feedback_message_id=None, mr_inter_iteration_countdown_task_ref=None,
        mr_current_feedback_revert_task_ref=None, mr_triggering_event_for_menu=msg_ctx,
    )
    instruction_text = ("<b>Тест умственного вращения</b>\n\n"
                        "Вам будет показан 3D объект и 4 варианта 2D проекций. "
                        "Выберите номер той проекции, через которую мог бы пройти 3D объект.\n"
                        "Объект можно вращать мысленно перед 'просовыванием'.\n\n"
                        f"Тест состоит из {MENTAL_ROTATION_NUM_ITERATIONS} заданий.")
    kbd = InlineKeyboardMarkup(inline_keyboard=[[IKB(text="Начать тест", callback_data="mr_ack_instructions")]])

    if not app_settings.MR_REFERENCE_FILES: # Use app_settings.
        logger.error("MR Start: CRITICAL - settings.MR_REFERENCE_FILES is empty. Test cannot start.")
        await bot_instance.send_message(chat_id, "Критическая ошибка: Пул эталонных изображений для теста пуст. Тест не может быть запущен. Пожалуйста, сообщите администратору.")
        await _safe_delete_message(bot_instance, chat_id, (await state.get_data()).get("status_message_id_to_delete_later"), "MR start common status error")
        await state.clear(); return

    try:
        await bot_instance.send_message(chat_id, instruction_text, reply_markup=kbd, parse_mode=ParseMode.HTML)
        if isinstance(trigger_event, CallbackQuery) and trigger_event.message: await trigger_event.message.delete() # Delete menu message
    except Exception as e_start_instr:
        logger.error(f"MR start: Error sending initial instructions: {e_start_instr}", exc_info=True)
        await bot_instance.send_message(chat_id, "Ошибка при запуске теста умственного вращения. Попробуйте /start.")
        await _safe_delete_message(bot_instance, chat_id, (await state.get_data()).get("status_message_id_to_delete_later"), "MR start common status fail")
        await state.clear()


@router.callback_query(F.data == "mr_ack_instructions", MentalRotationStates.initial_instructions_mr)
async def mr_ack_instructions_callback(cb: CallbackQuery, state: FSMContext, bot: Bot):
    await cb.answer()
    await state.update_data(mr_test_start_time=time.time())
    if cb.message: await cb.message.delete()
    chat_id = (await state.get_data()).get("mr_chat_id")
    if chat_id: await _display_mr_stimulus(chat_id, state, bot, trigger_msg_context=cb.message)
    else:
        logger.error("MR Ack Instr: chat_id missing. Aborting."); 
        if cb.message: await cb.message.answer("Ошибка: ID чата не найден. Тест прерван.")
        await _finish_mental_rotation_test(state, bot, None, True, True, trigger_msg_context=cb.message)


@router.callback_query(F.data.startswith("mr_answer_"), MentalRotationStates.displaying_stimulus_mr)
async def mr_answer_callback(cb: CallbackQuery, state: FSMContext, bot: Bot):
    await cb.answer()
    data = await state.get_data()
    chat_id = data.get("mr_chat_id")
    if not chat_id: # Should not happen if test started correctly
        logger.error("MR Answer: chat_id missing. Aborting."); 
        if cb.message: await cb.message.answer("Критическая ошибка: ID чата не найден. Тест прерван.")
        await _finish_mental_rotation_test(state, bot, None, True, True, trigger_msg_context=cb.message); return

    reaction_time_s = round(time.time() - data.get("mr_iteration_start_time", time.time()), 2)
    selected_option_idx = int(cb.data.split("_")[-1]) - 1
    correct_option_idx = data.get("mr_correct_option_index_for_current_iter")
    is_correct = selected_option_idx == correct_option_idx

    iteration_data = {"iteration": data.get("mr_current_iteration"), "is_correct": is_correct, "reaction_time_s": reaction_time_s, "selected_option": selected_option_idx + 1, "correct_option": (correct_option_idx + 1 if correct_option_idx is not None else "N/A")}
    current_results = data.get("mr_iteration_results", []); current_results.append(iteration_data)
    await state.update_data(mr_iteration_results=current_results)

    feedback_text_bold, feedback_text_normal = f"<b>{'Верно!' if is_correct else 'Неверно!'}</b>", f"{'Верно!' if is_correct else 'Неверно!'}"
    feedback_msg_id = data.get("mr_feedback_message_id")
    
    previous_revert_task = data.get("mr_current_feedback_revert_task_ref")
    if previous_revert_task and not previous_revert_task.done(): previous_revert_task.cancel(); await asyncio.sleep(0.01)

    try:
        if feedback_msg_id: await bot.edit_message_text(text=feedback_text_bold, chat_id=chat_id, message_id=feedback_msg_id, parse_mode=ParseMode.HTML)
        else: msg_fb = await bot.send_message(chat_id, feedback_text_bold, parse_mode=ParseMode.HTML); feedback_msg_id = msg_fb.message_id; await state.update_data(mr_feedback_message_id=feedback_msg_id)
        if feedback_msg_id:
            revert_task = asyncio.create_task(_mr_schedule_feedback_revert(chat_id, feedback_msg_id, feedback_text_normal, bot, state))
            await state.update_data(mr_current_feedback_revert_task_ref=revert_task)
    except Exception as e_fb_ans: logger.error(f"MR Answer: Feedback msg error: {e_fb_ans}", exc_info=True)
    
    options_msg_id = data.get("mr_options_message_id")
    if options_msg_id and chat_id:
        try: await bot.edit_message_reply_markup(chat_id=chat_id, message_id=options_msg_id, reply_markup=None)
        except: pass # Ignore if already removed or error

    await state.set_state(MentalRotationStates.processing_answer_mr)
    await _mr_proceed_to_next_iteration_or_finish(state, bot, chat_id, trigger_msg_context=cb.message)


async def save_mental_rotation_results(
    trigger_msg_context: Optional[Message], # Can be None if called internally without direct message trigger
    state: FSMContext, 
    is_interrupted: bool = False,
    target_sheet_name: Optional[str] = None, # New parameter
):
    logger.info(f"Saving Mental Rotation results. Interrupted: {is_interrupted}. Sheet: {target_sheet_name or 'default'}")
    data = await state.get_data()
    
    uid = data.get("original_user_uid") or data.get("mr_unique_id_for_test") or data.get("active_unique_id")
    profile_info = await get_active_profile_from_fsm(state)
    p_name, p_age, p_tgid = "", "", ""

    if profile_info and profile_info.get('unique_id') == uid:
        p_name, p_age, p_tgid = profile_info.get("name"), profile_info.get("age"), profile_info.get("telegram_id")
    else: # Fallback to MR specific if active profile doesn't match or is missing
        p_name = data.get("mr_profile_name_for_test", data.get("active_name"))
        p_age = data.get("mr_profile_age_for_test", data.get("active_age"))
        p_tgid = data.get("mr_profile_telegram_id_for_test", data.get("active_telegram_id"))

    if not uid: logger.error("MR Save Results: UID not found. Cannot save results."); return

    correct_ans = data.get("mr_final_correct_answers", 0)
    avg_rt = data.get("mr_final_avg_reaction_time_s", 0.0)
    total_time = data.get("mr_final_total_test_time_s", 0.0)
    ind_resp_str = data.get("mr_final_individual_responses_str", "N/A")
    interrupted_status = "Да" if data.get("mr_final_interrupted_status", is_interrupted) else "Нет"

    from openpyxl import load_workbook # Local import
    try:
        if not os.path.exists(EXCEL_FILENAME): initialize_excel_file()
        wb = load_workbook(EXCEL_FILENAME)
        
        ws = None
        if target_sheet_name and target_sheet_name in wb.sheetnames: ws = wb[target_sheet_name]
        elif wb.sheetnames: ws = wb[wb.sheetnames[0]]; logger.warning(f"MR Save: Target sheet '{target_sheet_name}' not found/specified. Using '{ws.title}'.")
        else: logger.error(f"MR Save: No sheets in '{EXCEL_FILENAME}'."); return

        sheet_headers = [cell.value for cell in ws[1]] if ws.max_row > 0 else []
        if not sheet_headers or "Unique ID" not in sheet_headers:
            logger.error(f"MR Save: Sheet '{ws.title}' lacks headers or 'Unique ID'. Re-initializing.");
            initialize_excel_file(); wb = load_workbook(EXCEL_FILENAME)
            if target_sheet_name and target_sheet_name in wb.sheetnames: ws = wb[target_sheet_name]
            else: ws = wb[wb.sheetnames[0]] # Fallback to first sheet
            sheet_headers = [cell.value for cell in ws[1]]
            if not sheet_headers or "Unique ID" not in sheet_headers:
                 raise ValueError(f"Sheet '{ws.title}' still lacks 'Unique ID' after re-init.")
        
        header_to_idx_map = {header: idx for idx, header in enumerate(sheet_headers)}
        uid_col_idx = header_to_idx_map.get("Unique ID")
        
        row_num = -1
        if uid_col_idx is not None:
            for r_idx, row_vals_tuple in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                if len(row_vals_tuple) > uid_col_idx and str(row_vals_tuple[uid_col_idx]) == str(uid):
                    row_num = r_idx; break
        
        if row_num == -1:
            new_row_data = [""] * len(sheet_headers)
            if uid_col_idx is not None: new_row_data[uid_col_idx] = uid
            ws.append(new_row_data); row_num = ws.max_row
            logger.info(f"MR Save: New row for UID {uid} in sheet '{ws.title}' at row {row_num}.")

        # Populate BASE_HEADERS & MENTAL_ROTATION_HEADERS
        for header_key in BASE_HEADERS + MENTAL_ROTATION_HEADERS:
            if header_key in header_to_idx_map:
                col_idx = header_to_idx_map[header_key]
                current_val_in_sheet = ws.cell(row=row_num, column=col_idx + 1).value
                value_to_write = None

                if header_key == "Name": value_to_write = p_name
                elif header_key == "Age": value_to_write = p_age
                elif header_key == "Telegram ID": value_to_write = p_tgid
                elif header_key == "Unique ID": value_to_write = uid
                elif header_key == "MentalRotation_CorrectAnswers": value_to_write = correct_ans
                elif header_key == "MentalRotation_AverageReactionTime_s": value_to_write = avg_rt
                elif header_key == "MentalRotation_TotalTime_s": value_to_write = total_time
                elif header_key == "MentalRotation_IndividualResponses": value_to_write = ind_resp_str
                elif header_key == "MentalRotation_Interrupted": value_to_write = interrupted_status
                
                if value_to_write is not None and (current_val_in_sheet is None or str(current_val_in_sheet) != str(value_to_write)):
                     ws.cell(row=row_num, column=col_idx + 1).value = value_to_write
            else:
                logger.warning(f"MR Save: Header '{header_key}' not found in sheet '{ws.title}'.")

        wb.save(EXCEL_FILENAME)
        logger.info(f"Mental Rotation results for UID {uid} saved to sheet '{ws.title}'. Interrupted: {interrupted_status}")
    except Exception as e_save_excel_mr:
        logger.error(f"MR Save Results: Excel save error for UID {uid}, Sheet '{target_sheet_name or 'default'}': {e_save_excel_mr}", exc_info=True)
        # Avoid sending message if trigger_msg_context is None (e.g. called from _finish_mental_rotation_test without user interaction)
        if trigger_msg_context and hasattr(trigger_msg_context, 'chat') and await state.get_state() is not None:
            try: await trigger_msg_context.answer("Ошибка при сохранении результатов Теста умственного вращения.")
            except: pass # Best effort


async def cleanup_mental_rotation_ui(state: FSMContext, bot_instance: Bot, final_text: str | None = None):
    data = await state.get_data()
    chat_id = data.get("mr_chat_id")
    logger.info(f"MR Cleanup UI: Chat {chat_id if chat_id else 'N/A'}. Final text: '{final_text}'")

    for task_key in ["mr_inter_iteration_countdown_task_ref", "mr_current_feedback_revert_task_ref"]:
        task = data.get(task_key)
        if task and not task.done(): task.cancel(); await asyncio.sleep(0.02)
    
    mr_ui_msg_ids_to_clean = {data.get(k) for k in ["mr_reference_message_id", "mr_options_message_id", "mr_countdown_message_id", "mr_feedback_message_id"] if data.get(k)}
    
    last_msg_id_for_edit = None
    if final_text and chat_id: # Determine which message to edit if final_text is provided
        for key in reversed(["mr_options_message_id", "mr_reference_message_id", "mr_feedback_message_id", "mr_countdown_message_id"]):
            if data.get(key): last_msg_id_for_edit = data.get(key); break
    
    if chat_id:
        if final_text and last_msg_id_for_edit:
            is_photo = last_msg_id_for_edit in [data.get("mr_options_message_id"), data.get("mr_reference_message_id")]
            try:
                if is_photo: await bot_instance.edit_message_caption(chat_id=chat_id, message_id=last_msg_id_for_edit, caption=final_text, reply_markup=None)
                else: await bot_instance.edit_message_text(text=final_text, chat_id=chat_id, message_id=last_msg_id_for_edit, reply_markup=None)
                mr_ui_msg_ids_to_clean.discard(last_msg_id_for_edit)
            except Exception: # If edit fails, send new and delete all old
                try: await bot_instance.send_message(chat_id, final_text)
                except: pass 
        elif final_text: # No specific message to edit, but final_text needs to be sent
             try: await bot_instance.send_message(chat_id, final_text)
             except: pass

        for msg_id_del in mr_ui_msg_ids_to_clean: # Delete remaining messages
            try: await bot_instance.delete_message(chat_id, msg_id_del)
            except: pass 

    fsm_keys_to_remove = [key for key in data if key.startswith("mr_")]
    for key_to_remove in fsm_keys_to_remove: del data[key_to_remove]
    await state.set_data(data)
    logger.info(f"MR Cleanup UI: FSM data cleaned of mr_ prefixed keys. Chat: {chat_id or 'N/A'}")
